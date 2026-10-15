import pytest
from unittest.mock import MagicMock, patch 
from openpyxl import Workbook, workbook 
from main.service.vacation_manager import save_vacation_request, check_vacation_request_availability, find_employee_row, load_vacation_excel, unsuccessful_requests

# ============================================================ 
# # find_employee_row() # 
# ============================================================
def test_find_employee_row_existing_employee(): 
        workbook = Workbook() 
        sheet = workbook.active 

        sheet.append(["Employee", "1", "2", "3"]) 
        sheet.append(["Dimitris Vampiris", "", "", ""]) 
        sheet.append(["John Doe", "", "", ""]) 

        row_number, employee_row = find_employee_row( sheet, "Dimitris Vampiris" ) 

        assert row_number == 2 
        assert employee_row[0] == "Dimitris Vampiris" 

def test_find_employee_row_case_and_whitespace(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.append(["Employee", "1", "2", "3"]) 
    sheet.append(["Dimitris Vampiris", "", "", ""]) 

    row_number, employee_row = find_employee_row( sheet, " DIMITRIS VAMPIRIS " ) 

    assert row_number == 2 
    assert employee_row[0] == "Dimitris Vampiris" 

def test_find_employee_row_employee_not_found(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.append(["Employee", "1", "2", "3"]) 
    sheet.append(["Dimitris Vampiris", "", "", ""]) 

    row_number, employee_row = find_employee_row( sheet, "John Doe" ) 

    assert row_number is None 
    assert employee_row is None


# ============================================================ 
# # check_vacation_request_availability() # 
# ============================================================

def test_check_vacation_available(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.title = "November" 
    # Column A = employee 
    # # Column B = day 1 
    # # Column C = day 2 
    # # ... 
    sheet.append(["Employee", "1", "2", "3", "4", "5"]) 
    sheet.append(["Dimitris Vampiris", "", "", "", "", ""]) 
    sheet.append(["John Doe", "", "", "", "", ""]) 

    result = check_vacation_request_availability( sheet, "Dimitris Vampiris", 1, 3 ) 

    assert result is True 
    assert unsuccessful_requests == []


def test_check_vacation_employee_not_found(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.title = "November" 
    sheet.append(["Employee", "1", "2", "3", "4", "5"]) 
    sheet.append(["Dimitris Vampiris", "", "", "", "", ""]) 

    result = check_vacation_request_availability( sheet, "John Doe", 1, 3 ) 

    assert result is False 
    assert unsuccessful_requests == [ "Employee John Doe not found in the vacation excel file." ]

def test_check_vacation_already_requested(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.title = "November" 
    sheet.append(["Employee", "1", "2", "3", "4", "5"]) 
    # Day 2 already has X for Dimitris 
    sheet.append(["Dimitris Vampiris", "", "X", "", "", ""]) 

    result = check_vacation_request_availability( sheet, "Dimitris Vampiris", 1, 3 ) 

    assert result is False 
    assert unsuccessful_requests == [ "You have already requested vacation for 2 of November." ]

def test_check_vacation_two_other_employees_already_on_vacation(): 
    workbook = Workbook() 
    sheet = workbook.active 

    sheet.title = "November" 
    sheet.append(["Employee", "1", "2", "3", "4", "5"]) 
    # Two employees already have vacation on day 2 
    sheet.append(["Employee 1", "", "X", "", "", ""]) 
    sheet.append(["Employee 2", "", "X", "", "", ""]) 
    sheet.append(["Dimitris Vampiris", "", "", "", "", ""]) 

    result = check_vacation_request_availability( sheet, "Dimitris Vampiris", 1, 3 ) 

    assert result is False 
    assert unsuccessful_requests == [ "You cannot request vacation for 2 of November as there are already 2 employees on vacation." ]

# ============================================================ 
# # load_vacation_excel() # 
# ============================================================

@patch("main.service.vacation_manager.load_workbook") 
def test_load_vacation_excel(mock_load_workbook): 
    workbook = MagicMock() 
    sheet = MagicMock() 

    workbook.__getitem__.return_value = sheet 
    mock_load_workbook.return_value = workbook 
    result_workbook, result_sheet = load_vacation_excel("November") 

    mock_load_workbook.assert_called_once_with( "main/resources/documents/Vacation Planner.xlsx" ) 

    workbook.__getitem__.assert_called_once_with("November") 

    assert result_workbook == workbook 
    assert result_sheet == sheet

# ============================================================ 
# # save_vacation_request() # 
# ============================================================

@patch("main.service.vacation_manager.check_vacation_request_availability") 
@patch("main.service.vacation_manager.load_vacation_excel") 
def test_save_vacation_request_success( mock_load_vacation_excel, mock_check_availability ): 
    workbook = MagicMock() 
    sheet = MagicMock() 

    mock_load_vacation_excel.return_value = ( workbook, sheet ) 
    mock_check_availability.return_value = True 

    # find_employee_row is called internally. 
    # Mock it so we don't depend on a real Excel sheet. 
    with patch( 
        "main.service.vacation_manager.find_employee_row", 
        return_value=(5, MagicMock()) 
    ) as mock_find_employee_row: 
        result = save_vacation_request( "November", "Dimitris Vampiris", 28, 30 ) 

    assert result == ( "Dimitris Vampiris, " "I have saved successfully your vacation request." ) 

    mock_load_vacation_excel.assert_called_once_with("November") 
    mock_check_availability.assert_called_once_with( sheet, "Dimitris Vampiris", 28, 30 ) 

    assert sheet.cell.call_count == 3 

@patch("main.service.vacation_manager.check_vacation_request_availability") 
@patch("main.service.vacation_manager.load_vacation_excel") 
def test_save_vacation_request_unavailable( mock_load_vacation_excel, mock_check_availability ): 
    workbook = MagicMock() 
    sheet = MagicMock() 
    mock_load_vacation_excel.return_value = ( workbook, sheet ) 
    mock_check_availability.return_value = False 

    unsuccessful_requests.clear() 
    unsuccessful_requests.append( "You have already requested vacation for 28 of November." ) 

    result = save_vacation_request( "November", "Dimitris Vampiris", 28, 30 ) 

    assert result == ( "You have already requested vacation for 28 of November." ) 
    # Excel must NOT be modified when the request is unavailable. 
    sheet.cell.assert_not_called() 

    # Excel must NOT be saved. 
    workbook.save.assert_not_called()